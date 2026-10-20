"""Inspect pinned USDA 2016 maize daily records; never run a model or fill gaps.

Requires openpyxl from requirements.txt. Run from the checkout with
python -m scripts.audit_bushland_data --output /tmp/bushland-screening.json.
Only the daily water-balance tab is summarized; other tabs remain source evidence.
"""

import argparse
import hashlib
import json
import math
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'data/field_reference/bushland_2016'
FILES = {'2016_Maize_water_balance.xlsx', 'Bushland_TX_soil_properties.xlsx',
         'Conventions for Bushland, TX, Weighing Lysimeter Datasets.xlsx'}
SHEET = '2016 Maize Daily WBal'
LOCATIONS = ('NE', 'SE', 'NW', 'SW')
QUANTITIES = ('ET in mm', 'dew in mm', 'precip in mm', 'irrig. In mm',
              'CW adjust. In mm', 'drainage in mm')
HEADERS = ('Year', 'Day of year') + tuple(
    f'{location} {quantity}' for location in LOCATIONS for quantity in QUANTITIES)


def summarize(rows):
    """Keep blanks, declared missing markers and signed source values distinct."""
    if not rows or tuple(rows[0]) != HEADERS:
        raise ValueError('daily water-balance columns differ from the inspected source')
    data = rows[1:]
    if len(data) != 366:
        raise ValueError('expected all 366 source calendar rows, including missing values')
    for day, row in enumerate(data, 1):
        if (len(row) != len(HEADERS) or type(row[0]) is not int or row[0] != 2016
                or type(row[1]) is not int or row[1] != day):
            raise ValueError('source year/day sequence or row width changed')
    columns = {}
    for index, name in enumerate(HEADERS[2:], 2):
        values, marker_days, blank_days = [], [], []
        for row in data:
            value = row[index]
            if value is None:
                blank_days.append(row[1])
            elif value == '#N/A':
                marker_days.append(row[1])
            elif type(value) in (int, float) and math.isfinite(value):
                values.append(value)
            else:
                raise ValueError(f'unrecognized value at day {row[1]}, {name}')
        columns[name] = dict(numeric_count=len(values), missing_marker_days=marker_days,
            blank_days=blank_days, negative_count=sum(value < 0 for value in values),
            zero_count=sum(value == 0 for value in values),
            minimum=min(values) if values else None, maximum=max(values) if values else None)
    complete = {}
    for location in LOCATIONS:
        indices = [HEADERS.index(f'{location} {quantity}') for quantity in QUANTITIES]
        complete[location] = sum(all(type(row[index]) in (int, float)
                                    for index in indices) for row in data)
    return dict(sheet=SHEET, rows=len(data), columns=len(HEADERS), year=2016,
                day_of_year_first=1, day_of_year_last=366,
                columns_observed=columns, rows_with_all_six_numeric_values=complete)


def audit(source=SOURCE):
    manifest = json.loads((source / 'manifest.json').read_text(encoding='utf-8'))
    if manifest['schema_version'] != 1:
        raise ValueError('unsupported source manifest version')
    if (len(manifest['files']) != len(FILES)
            or {item['name'] for item in manifest['files']} != FILES):
        raise ValueError('source manifest must pin each of the three inspected workbooks')
    checked = []
    for item in manifest['files']:
        path = (source / item['name']).resolve()
        path.relative_to(source.resolve())
        raw = path.read_bytes()
        if (len(raw) != item['bytes'] or hashlib.md5(raw).hexdigest() != item['repository_md5']
                or hashlib.sha256(raw).hexdigest() != item['sha256']):
            raise ValueError(f"source checksum mismatch: {item['name']}")
        checked.append(item)
    path = source / '2016_Maize_water_balance.xlsx'
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        sheet = workbook[SHEET]
        if sheet.max_column != len(HEADERS) or sheet.max_row != 367:
            raise ValueError('daily sheet dimensions differ from the inspected source')
        summary = summarize(list(sheet.values))
    finally:
        workbook.close()
    formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        formula_cells = [cell.coordinate for row in formulas[SHEET] for cell in row
                         if cell.data_type == 'f']
    finally:
        formulas.close()
    return dict(schema_version=1, audited_on=manifest['audited_on'],
        scope=manifest['scope'], license=manifest['license'],
        publication_note=manifest['publication_note'], checked_files=checked,
        observed_file_summary=summary, daily_formula_cells=formula_cells,
        parsing='openpyxl reads stored numeric values and formula caches; no formula recalculation, '
                'author code execution, missing-value replacement, timestamp conversion or unit conversion.',
        reviewer_assessed_limitations=[
            'Source time uses fixed US Central Standard Time, without daylight-saving adjustment; '
            'subdaily timestamps denote interval ends. No timezone conversion was performed.',
            'Lysimeter relative water storage is not absolute soil profile water storage.',
            'The drainage columns describe drainage-tank emptying, not daily soil drainage flux. '
            'Counterweight adjustments are instrument operations, not water applications.',
            'Annual daily rows include non-growing periods. A crop calendar and matching weather, '
            'initial soil profile observations and quality review remain necessary.',
            'NE/SE and NW/SW records must retain their site and irrigation-treatment identities. '
            'No independence or train/test split is asserted.',
            'The soil workbook includes nonnumeric clay text (43.S) and site-specific hydraulic '
            'behavior; no soil parameters were imported or corrected.',
            'Missing markers and blanks remain missing. '
            'A numeric cell alone does not establish observation quality.',
            'Texas records do not establish performance in Burkina Faso or Taiwan.'],
        model_validation_performed=False)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path)
    args = parser.parse_args()
    result = json.dumps(audit(), indent=2, ensure_ascii=False, allow_nan=False) + '\n'
    if args.output:
        args.output.write_text(result, encoding='utf-8')
    else:
        print(result, end='')


if __name__ == '__main__':
    main()
