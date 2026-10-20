"""Prepare dated USDA weather summaries and compare paired rainfall representations.

Offline source assessment, not model fitting, ET validation or farm-account ingestion.
Requires openpyxl. Original missing values, signed rain and untimed cells are retained.
"""

import argparse
import hashlib
import json
import math
from datetime import date, datetime, time, timedelta

from openpyxl import load_workbook

from scripts.audit_bushland_data import SOURCE
from scripts.audit_bushland_data import audit as audit_balance

WEATHER_FILE = '2016_15-min_weather_SWMRU_CPRL_1.xlsx'
WEATHER_SHEET = '2016 15-min weather'
HEADERS = ('Date time', 'Decimal day of year', 'Year', 'Day of year', 'Time in hhmm',
    'Solar irradiance in W/m2', 'RH in %', 'Air temperature in °C', 'Wind speed in m/s',
    'Air pressure in kPa', 'Average precipitation in mm',
    'NE lysimeter precipitaton in mm', 'SE lysimeter precipitaton in mm',
    'NW lysimeter precipitaton in mm', 'SW lysimeter precipitaton in mm')
LOCATIONS = ('NE', 'SE', 'NW', 'SW')
# Comparison tolerance for summation roundoff, not measurement accuracy or model acceptance.
RAIN_COMPARISON_TOLERANCE_MM = 1e-8


def _value(value):
    if value is None or value == '#N/A':
        return None
    if type(value) not in (int, float) or not math.isfinite(value):
        raise ValueError(f'unrecognized numeric cell: {value!r}')
    return value


def summarize_day(rows, day):
    if type(day) is not date or day.year != 2016 or len(rows) != 96:
        raise ValueError('supply exactly 96 consecutive interval endpoints for a 2016 day')
    doy = day.timetuple().tm_yday
    for interval, row in enumerate(rows, 1):
        endpoint = datetime.combine(day, time.min) + timedelta(minutes=15*interval)
        hhmm = 2400 if interval == 96 else endpoint.hour*100+endpoint.minute
        if (len(row) != 15 or row[0] != endpoint or type(row[0]) is not datetime
                or row[0].tzinfo is not None or row[2:5] != (2016, doy, hhmm)
                or any(type(v) is not int for v in row[2:5])
                or type(row[1]) not in (int, float) or not math.isfinite(row[1])
                or not math.isclose(row[1], doy+interval/96, abs_tol=1e-8, rel_tol=0)):
            raise ValueError(f'inconsistent interval time or source day at {day}, slot {interval}')
    columns = [[_value(row[column]) for row in rows] for column in range(5, 15)]

    def complete(index, aggregate):
        values = columns[index]
        return aggregate(values) if all(v is not None for v in values) else None

    precipitation = {}
    for index, location in enumerate(('four_lysimeter_average',) + LOCATIONS, 5):
        values = columns[index]
        precipitation[location] = dict(
            total_mm=complete(index, math.fsum),
            missing_intervals=sum(v is None for v in values),
            negative_intervals=sum(v is not None and v < 0 for v in values))
    return dict(day=day.isoformat(), source_day_of_year=doy, intervals=96,
        solar_shortwave_mj_m2=complete(0, lambda values: math.fsum(values)*900/1e6),
        minimum_15min_mean_rh_pct=complete(1, min),
        maximum_15min_mean_rh_pct=complete(1, max),
        minimum_15min_mean_temperature_c=complete(2, min),
        maximum_15min_mean_temperature_c=complete(2, max),
        mean_2m_wind_m_s=complete(3, lambda values: math.fsum(values)/96),
        mean_pressure_kpa=complete(4, lambda values: math.fsum(values)/96),
        missing_meteorology_intervals=dict(zip(HEADERS[5:10],
            (sum(v is None for v in values) for values in columns[:5]))),
        precipitation=precipitation)


def _daily_rows(sheet, width):
    rows = list(sheet.iter_rows(max_col=width, values_only=True))
    if len(rows) != 367:
        raise ValueError('daily source must contain its header and 366 calendar rows')
    return rows


def prepare(source=SOURCE):
    manifest = json.loads((source/'weather_manifest.json').read_text(encoding='utf-8'))
    item = manifest['file']
    if manifest['schema_version'] != 1 or item['name'] != WEATHER_FILE:
        raise ValueError('unexpected weather manifest')
    raw = (source/WEATHER_FILE).read_bytes()
    if (len(raw) != item['bytes'] or hashlib.md5(raw).hexdigest() != item['repository_md5']
            or hashlib.sha256(raw).hexdigest() != item['sha256']):
        raise ValueError('weather source checksum mismatch')
    balance_audit = audit_balance(source)
    workbook = load_workbook(source/WEATHER_FILE, read_only=True, data_only=True, keep_links=False)
    try:
        rows = list(workbook[WEATHER_SHEET].values)
        if rows[0] != HEADERS or len(rows) != 35138:
            raise ValueError('unexpected source weather shape or headers')
        # The pinned source has two zero cells after the last dated row. Expose them;
        # never silently drop arbitrary untimed data or attribute it to December 31.
        if rows[-1] != (None,)*11 + (0, 0, None, None):
            raise ValueError('untimed trailing cells differ from the inspected source')
        daily = _daily_rows(workbook['2016 daily precipitation'], 7)
        if daily[0] != ('Date', 'Day of year') + HEADERS[10:]:
            raise ValueError('daily precipitation headers changed')
    finally:
        workbook.close()
    workbook = load_workbook(source/'2016_Maize_water_balance.xlsx',
        read_only=True, data_only=True, keep_links=False)
    try:
        balance = _daily_rows(workbook['2016 Maize Daily WBal'], 26)
    finally:
        workbook.close()
    days, differences = [], []
    comparison_counts = {name: dict(available_pairs=0, missing_pairs=0, differences=0)
        for name in ('published_weather_daily_mm', 'published_balance_daily_mm')}
    for offset in range(366):
        day = date(2016, 1, 1) + timedelta(days=offset)
        record = summarize_day(rows[1+96*offset:1+96*(offset+1)], day)
        source_daily = daily[offset+1]
        if source_daily[:2] != (datetime.combine(day, time.min), offset+1):
            raise ValueError('daily precipitation calendar sequence changed')
        for location_index, location in enumerate(('four_lysimeter_average',)+LOCATIONS):
            rain = record['precipitation'][location]
            rain['published_weather_daily_mm'] = _value(source_daily[2+location_index])
            if location in LOCATIONS:
                start = 2+6*LOCATIONS.index(location)
                rain['published_balance_daily_mm'] = _value(balance[offset+1][start+2])
                rain['published_balance_et_mm'] = _value(balance[offset+1][start])
            for comparison in ('published_weather_daily_mm', 'published_balance_daily_mm'):
                if comparison not in rain:
                    continue
                other = rain.get(comparison)
                total = rain['total_mm']
                if total is not None and other is not None:
                    comparison_counts[comparison]['available_pairs'] += 1
                    delta = total-other
                    if abs(delta) > RAIN_COMPARISON_TOLERANCE_MM:
                        comparison_counts[comparison]['differences'] += 1
                        differences.append(dict(day=day.isoformat(), location=location,
                            comparison=comparison, interval_sum_mm=total,
                            published_daily_mm=other, difference_mm=delta))
                else:
                    comparison_counts[comparison]['missing_pairs'] += 1
        days.append(record)
    return dict(schema_version=1, source_weather=manifest,
        source_balance_files=balance_audit['checked_files'], timezone='Etc/GMT+6',
        interval_convention='Fixed UTC-06:00, without daylight-saving time. Each timestamp ends '
            'a 15-minute interval; 24:00 belongs to the source day that just ended.',
        aggregation='All 96 values required for each quantity. Solar shortwave is the sum of '
            '15-minute mean W/m2 times 900 seconds / 1e6. Temperature/RH extrema are extrema '
            'of interval means, not instantaneous daily extremes. No values are filled or clipped.',
        untimed_source_cells=dict(sheet=WEATHER_SHEET, row=35138,
            cells={'L35138': 0, 'M35138': 0}, handling='Retained here; excluded from dated summaries.'),
        calendar_days=len(days), dated_intervals=96*len(days),
        rain_comparison_tolerance_mm=RAIN_COMPARISON_TOLERANCE_MM,
        comparison_kind='Consistency of related published rainfall representations; '
            'not independent observations or validation of an Agronaut model.',
        comparison_counts=comparison_counts,
        rainfall_differences=differences,
        model_input_ready=False, model_validation_performed=False,
        missing_model_inputs=['Daily net radiation is absent; solar shortwave cannot replace it.',
            'Daily instantaneous temperature/humidity extremes are not supplied by these interval means.',
            'Crop calendar, applicable coefficients, initial soil state and an agreed evaluation protocol remain unassembled.'],
        days=days)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', required=True)
    args = parser.parse_args()
    from pathlib import Path
    Path(args.output).write_text(json.dumps(prepare(), indent=2, ensure_ascii=False,
        allow_nan=False)+'\n', encoding='utf-8')


if __name__ == '__main__':
    main()
