"""Compare pinned rainfall representations without choosing or changing model inputs.

Offline source consistency assessment, not calibration or irrigation validation.
Requires openpyxl; run python -m scripts.reconcile_bushland_rain --output FILE.
"""
import argparse
import json
import math
from datetime import date, timedelta
from pathlib import Path

from openpyxl import load_workbook

from scripts.prepare_bushland_weather import (
    RAIN_COMPARISON_TOLERANCE_MM,
    SOURCE,
    WEATHER_FILE,
    WEATHER_SHEET,
    _value,
    prepare,
)

BALANCE_FILE = '2016_Maize_water_balance.xlsx'
BALANCE_SHEET = '2016 Maize 15-min WBal'
HEADERS = ('Year', 'Decimal DOY', 'NE dew/frost in mm', 'SE dew/frost in mm',
    'NE precip in mm', 'SE precip in mm', 'NE irrig in mm', 'SE irrig in mm',
    'NE ET in mm', 'SE ET in mm', 'NW dew/frost in mm', 'SW dew/frost in mm',
    'NW precip in mm', 'SW precip in mm', 'NW irrig in mm', 'SW irrig in mm',
    'NW ET in mm', 'SW ET in mm')
POSITIONS = dict(NE=(11, 4, 'L', 'E'), SE=(12, 5, 'M', 'F'), NW=(13, 12, 'N', 'M'), SW=(14, 13, 'O', 'N'))


def compare_intervals(weather, balance):
    """Align declared interval endpoints; retain missing categories and exact numeric equality."""
    if len(weather) != 35136 or len(balance) != 35136:
        raise ValueError('expected 35136 intervals in each source')
    counts = {loc: dict(numeric_equal=0, missing_marker_equal=0, blank_equal=0, differences=0) for loc in POSITIONS}
    differences = []
    for index, (wx, wb) in enumerate(zip(weather, balance), 1):
        expected = 1+index/96
        if (len(wx) != 15 or len(wb) != 18 or type(wb[0]) is not int or wb[0] != 2016
                or type(wb[1]) not in (int, float) or not math.isfinite(wb[1])
                or not math.isclose(wb[1], expected, rel_tol=0, abs_tol=1e-8)
                or type(wx[1]) not in (int, float) or not math.isfinite(wx[1])
                or not math.isclose(wx[1], expected, rel_tol=0, abs_tol=1e-8)):
            raise ValueError(f'interval sequence changed at source row {index+1}')
        for location, (wi, bi, wc, bc) in POSITIONS.items():
            left, right = wx[wi], wb[bi]
            _value(left)
            _value(right)
            if left is None and right is None:
                key = 'blank_equal'
            elif left == right == '#N/A':
                key = 'missing_marker_equal'
            elif type(left) in (int, float) and type(right) in (int, float) and left == right:
                key = 'numeric_equal'
            else:
                key = 'differences'
                differences.append(dict(location=location, weather_cell=f'{wc}{index+1}', balance_cell=f'{bc}{index+1}',
                    weather_value=left, balance_value=right, interval_end_decimal_doy=expected))
            counts[location][key] += 1
    return dict(intervals_per_location=len(weather), counts=counts, differences=differences,
        comparison='Exact numeric equality; #N/A markers and blank cells remain separate categories.',
        time_alignment='Both decimal-DOY sequences must match 15-minute endpoints from 1+1/96 through 367 within 1e-8 day. '
            'This accepts stored floating-point drift, not a clock offset.')


def consecutive_groups(rows):
    """Group only adjacent differing days, retaining each signed residual and both totals."""
    groups = []
    for row in rows:
        day = date.fromisoformat(row['day'])
        if groups and day <= date.fromisoformat(groups[-1]['days'][-1]['day']):
            raise ValueError('daily differences must be strictly chronological')
        if not groups or day != date.fromisoformat(groups[-1]['days'][-1]['day'])+timedelta(days=1):
            groups.append(dict(days=[]))
        groups[-1]['days'].append(row)
    for group in groups:
        values = group['days']
        total = math.fsum(row['interval_sum_mm'] for row in values)
        published = math.fsum(row['published_balance_daily_mm'] for row in values)
        residual = published-total
        group.update(start=values[0]['day'], end=values[-1]['day'], day_count=len(values),
            interval_sum_mm=total, published_balance_sum_mm=published,
            published_minus_interval_mm=residual,
            totals_agree_within_tolerance=abs(residual) <= RAIN_COMPARISON_TOLERANCE_MM)
    return groups


def dew_comparison(rows, days, location, location_index):
    """Test the proposed universal dew explanation, without applying a correction."""
    if len(rows) != len(days):
        raise ValueError('daily dew and rainfall sequences must have equal lengths')
    result = dict(available_pairs=0, missing_pairs=0, differences_without_dew=0, differences_after_adding_dew=0)
    for row, day in zip(rows, days):
        rain = day['precipitation'][location]
        values = [rain['published_weather_daily_mm'], rain['published_balance_daily_mm'], _value(row[3+6*location_index])]
        if any(value is None for value in values):
            result['missing_pairs'] += 1
            continue
        weather, balance, dew = values
        result['available_pairs'] += 1
        result['differences_without_dew'] += int(abs(weather-balance) > RAIN_COMPARISON_TOLERANCE_MM)
        result['differences_after_adding_dew'] += int(abs(weather-(balance+dew)) > RAIN_COMPARISON_TOLERANCE_MM)
    return result


def reconcile(source=SOURCE):
    prepared = prepare(source)  # Verify pinned bytes, headers and complete dated sequence first.
    workbook = load_workbook(source/WEATHER_FILE, read_only=True, data_only=True, keep_links=False)
    try:
        weather = list(workbook[WEATHER_SHEET].iter_rows(min_row=2, max_row=35137, max_col=15, values_only=True))
    finally:
        workbook.close()
    workbook = load_workbook(source/BALANCE_FILE, read_only=True, data_only=True, keep_links=False)
    try:
        sheet = workbook[BALANCE_SHEET]
        rows = list(sheet.values)
        if sheet.max_row != 35137 or sheet.max_column != 18 or rows[0] != HEADERS:
            raise ValueError('water-balance interval headers or shape changed')
        intervals = compare_intervals(weather, rows[1:])
        daily = list(workbook['2016 Maize Daily WBal'].values)[1:]
    finally:
        workbook.close()
    locations = {}
    for index, location in enumerate(POSITIONS):
        differences = []
        for day in prepared['days']:
            rain = day['precipitation'][location]
            interval, published = rain['total_mm'], rain['published_balance_daily_mm']
            if interval is not None and published is not None and abs(published-interval) > RAIN_COMPARISON_TOLERANCE_MM:
                differences.append(dict(day=day['day'], interval_sum_mm=interval,
                    published_balance_daily_mm=published, published_minus_interval_mm=published-interval))
        groups = consecutive_groups(differences)
        locations[location] = dict(differing_days=len(differences), consecutive_groups=groups,
            groups_with_agreeing_totals=sum(group['totals_agree_within_tolerance'] for group in groups),
            daily_dew_hypothesis=dew_comparison(daily, prepared['days'], location, index))
    return dict(schema_version=1, source_weather=prepared['source_weather'],
        source_balance_files=prepared['source_balance_files'], tolerance_mm=RAIN_COMPARISON_TOLERANCE_MM,
        interval_comparison=intervals, locations=locations,
        weather_daily_comparison=prepared['comparison_counts']['published_weather_daily_mm'],
        model_input_ready=False, model_validation_performed=False, source_values_modified=False,
        interpretation=[
            'Interval equality establishes consistency of these related published representations, not independent replication or observation quality.',
            'Cancellation across consecutive differing days is consistent with different daily allocations. It does not prove the allocation rule, cause, or which daily series should drive a model.',
            'A universal dew-addition hypothesis is checked as weather daily minus (balance daily plus daily dew). No dew correction is applied.',
            'The separate weather daily table still differs from interval totals. No rainfall series is selected, shifted, filled or clipped.',
            'Temperature/humidity aggregation, radiation estimation, crop calendar, initial soil state, evaluation split and acceptance criteria remain unresolved before model evaluation.',
            'Texas source consistency does not establish performance in Burkina Faso or Taiwan.'],
        method_context=dict(url='https://acsess.onlinelibrary.wiley.com/doi/10.2134/age2018.09.0036',
            note='The published weather-quality methods distinguish precipitation from irrigation and discuss separate processing of lysimeter catch. They do not explain the specific discrepancies identified here.'))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', required=True, type=Path)
    args = parser.parse_args()
    args.output.write_text(json.dumps(reconcile(), indent=2, ensure_ascii=False, allow_nan=False)+'\n', encoding='utf-8')


if __name__ == '__main__':
    main()
